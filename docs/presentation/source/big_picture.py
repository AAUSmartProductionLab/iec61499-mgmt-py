"""The big picture of the talk on one slide: a vial travels through three stops (describe it, change it, let it
decide), each with a picture from the shape pack; a click moves it on and opens the next stop, and clicking a
stop opens it too. The pictures are drawn as a draw.io page (big-picture.drawio, one layer each), the rest
is native PowerPoint, so the slide can be pasted into a deck and edited.

    python big_picture.py <repo> <out.pptx>          (python-pptx; Chrome for the pictures)
"""
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).parent))
from drawlib import BLUE, GOLD, INK, PURPLE, Node, arrow, box, iso_shell, rect as d_rect, text as d_text, write  # noqa

REPO, OUT = Path(sys.argv[1]), Path(sys.argv[2])
FIG = REPO / "docs" / "figures"
DRAWIO = REPO / "docs" / "presentation" / "big-picture.drawio"
EXPORT = REPO / "docs" / "presentation" / "figures" / "big-picture"

# ==== The pictures, as a draw.io page ==================================================================
def icons(file, idx):
    return {c.get("id"): c for c in ET.parse(FIG / file).getroot().findall("diagram")[idx].find("mxGraphModel").iter("mxCell")}


L, C = icons("module-tools.drawio", 6), icons("module-tools.drawio", 4)


def place(node, cells, x, y, s=1.0):
    """Cells of a figure page, kept in their places relative to each other, scaled, their corner at (x, y)."""
    geos = [c.find("mxGeometry") for c in cells]
    x0, y0 = min(float(g.get("x")) for g in geos), min(float(g.get("y")) for g in geos)
    for c, g in zip(cells, geos):
        box(node, x + (float(g.get("x")) - x0) * s, y + (float(g.get("y")) - y0) * s, float(g.get("width")) * s,
            float(g.get("height")) * s, c.get("style"))


def shell_at(node, label, x, y, s=1.0):
    st, w, h, _ = iso_shell(label)
    box(node, x, y, w * s, h * s, st)


def badge(node, x, y, mark, fill, r=22):
    box(node, x - r, y - r, 2 * r, 2 * r, f"ellipse;whiteSpace=wrap;html=1;fillColor={fill};strokeColor=#FFFFFF;"
        f"strokeWidth=3;fontFamily=Arial;fontSize=22;fontColor=#FFFFFF;fontStyle=1;", mark)


h1, h2, h3, puck = Node("picture 1 · describe it"), Node("picture 2 · change it"), Node("picture 3 · let it decide"), \
    Node("the vial")
g = h1.group("A module: station, program, AAS")
place(g, [L["c104"], L["c105"]], 110, 220, 1.0)
place(g, [C["c5"]], 230, 120, 0.95)
shell_at(g, "Module", 400, 250, 0.9)
g = h2.group("modsync changes the program from the AAS")
place(g, [L["c104"], L["c105"]], 620, 230, 1.15)
shell_at(g, "Resource", 860, 210, 1.0)
arrow(g, [(900, 214), (840, 150), (745, 214)], GOLD, 5, label="modsync", size=24, label_color=GOLD)
g.items[-1]["xml"] = (lambda f: lambda p, ox, oy: f(p, ox, oy).replace("rounded=1;", "curved=1;rounded=1;"))(
    g.items[-1]["xml"])
g = h3.group("Agents: a product calls, modules bid")
place(g, [L["c5"]], 1250, 130, 0.62)
place(g, [L["c5"]], 1370, 200, 0.62)
place(g, [C["c18"]], 1150, 330, 1.5)
badge(g, 1160, 320, "P", PURPLE)
badge(g, 1300, 130, "R", BLUE)
badge(g, 1420, 200, "R", BLUE)
arrow(g, [(1186, 310), (1276, 145)], PURPLE, 3, True)
arrow(g, [(1186, 322), (1396, 210)], PURPLE, 3, True)
arrow(g, [(1290, 154), (1180, 300)], BLUE, 3)
place(puck.group("A vial on its puck"), [C["c18"]], 60, 600, 1.2)
write(DRAWIO, [("00 The big picture", 1600, 900, [h1, h2, h3, puck])])
subprocess.run([sys.executable, str(FIG / "export_layers.py"), str(DRAWIO), str(EXPORT), "--scale", "2"], check=True,
               capture_output=True)


def picture(slug):
    f = next(EXPORT.glob(f"[0-9][0-9]-{slug}*.png"))
    im = Image.open(f)
    bb = im.getchannel("A").getbbox()
    out = EXPORT / f"crop-{slug}.png"
    im.crop(bb).save(out)
    return out, (bb[2] - bb[0]) / (bb[3] - bb[1])


# ==== The slide ======================================================================================
PX = 7620
TEAL, TEAL_D, INK_T, GREY_T, LINE_T, LIGHT = "2A918B", "0F726C", "2D3134", "5C6166", "BEC2CA", "E6F3F2"
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes
E = lambda v: Emu(int(round(v * PX)))
rgb = RGBColor.from_string


def textbox(x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=0):
    tb = S.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for k, runs in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        if space and k:
            p.space_before = Pt(space)
        for text, pt, bold, colour, *it in runs:
            r = p.add_run()
            r.text = text
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(pt), bold, rgb(colour)
            r.font.italic = bool(it and it[0])
    return tb


def shape(kind, x, y, w, h, fill=None, line=None, width=1.5, alpha=None):
    sp = S.add_shape(kind, E(x), E(y), E(w), E(h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = sp.fill._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
            etree.SubElement(clr, "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha").set("val", str(alpha))
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb, sp.line.width = rgb(line), Pt(width)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def number(sp, n, colour):
    tf = sp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(22), True, rgb(colour)


textbox(85, 40, 900, 24, [[("OVERVIEW", 12, True, TEAL_D)]])
textbox(85, 70, 1430, 60, [[("Today: from a module to a line that changes itself", 28, True, TEAL)]],
        anchor=MSO_ANCHOR.MIDDLE)
textbox(85, 132, 1430, 30, [[("Three steps, each building on the one before. Click a step to open it.", 14, False,
                              GREY_T)]])

TRACK_Y, XS = 450, [300, 800, 1300]
track = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 110, TRACK_Y - 5, 1380, 10, "D9DCDF")
track.adjustments[0] = 0.5
STOPS = [
    ("1 · MODELLING", "Describe it", "What is a module, what describes it?",
     ["A module ships with its parts, its program and its AAS",
      "Products, processes and resources are AASs too",
      "Rules turn a skill in the AAS into function blocks"]),
    ("2 · RECONFIGURATION", "Change it", "Can a product change the program?",
     ["A step the line cannot do yet: what is missing?",
      "A new skill from the module's own primitives",
      "modsync builds it online and checks it"]),
    ("3 · NEXT: AGENTS", "Let it decide", "Can the line organise itself?",
     ["Products call for proposals, modules bid",
      "A module's agent reconfigures its module when it must",
      "The same tools, run by agents"]),
]
stops = []
for k, (x, (label, title, question, lines)) in enumerate(zip(XS, STOPS)):
    img, ratio = picture(f"picture-{k + 1}")
    ph = 215
    pw = min(380, ph * ratio)
    ph = pw / ratio
    pic = S.add_picture(str(img), E(x - pw / 2), E(418 - ph), E(pw), E(ph))
    pic.name = f"Picture {title}"
    dim = shape(MSO_SHAPE.RECTANGLE, x - pw / 2 - 4, 418 - ph - 4, pw + 8, ph + 8, "FFFFFF", alpha=35000)
    dim.name = f"Dim {title}"
    base = shape(MSO_SHAPE.OVAL, x - 32, TRACK_Y - 32, 64, 64, "FFFFFF", LINE_T, 2.5)
    number(base, k + 1, GREY_T)
    base.name = f"Stop {k + 1}"
    lit = shape(MSO_SHAPE.OVAL, x - 32, TRACK_Y - 32, 64, 64, TEAL, TEAL_D, 2.5)
    number(lit, k + 1, "FFFFFF")
    lit.name = f"Stop {k + 1} reached"
    head = textbox(x - 235, 494, 470, 80, [[(label, 12, True, TEAL_D)], [(title, 22, True, INK_T)]],
                   PP_ALIGN.CENTER, space=2)
    q = textbox(x - 235, 566, 470, 30, [[(question, 14, False, GREY_T, True)]], PP_ALIGN.CENTER)
    card = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x - 230, 606, 460, 206, LIGHT, TEAL, 1.5)
    card.adjustments[0] = 0.08
    card.name = f"Opened {title}"
    body = textbox(x - 210, 620, 420, 180, [[("›  " + s, 13, False, INK_T)] for s in lines], space=5)
    body.name = f"Opened {title} text"
    stops.append(dict(pic=pic, dim=dim, base=base, lit=lit, card=card, body=body))
img, ratio = picture("the-vial")
vial = S.add_picture(str(img), E(110 - 40), E(TRACK_Y - 64), E(80 * 1.0), E(80 / ratio))
vial.name = "The vial"
banner = shape(MSO_SHAPE.ROUNDED_RECTANGLE, 85, 832, 1430, 48, TEAL, None)
banner.adjustments[0] = 0.5
tf = banner.text_frame
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
for text, bold in [("One thread: ", True), ("the AAS describes, IEC 61499 runs, and tools, later agents, change it while "
                                          "it runs", False)]:
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(16), bold, rgb("FFFFFF")
banner.name = "The thread"

# ==== The animation: the main sequence, and a trigger on each stop ========================================
_id = [2]


def nid():
    _id[0] += 1
    return _id[0]


tgt = lambda s: f'<p:tgtEl><p:spTgt spid="{s.shape_id}"/></p:tgtEl>'


def fade(s, out=False, node="withEffect", dur=500):
    cls, tr = ("exit", "out") if out else ("entr", "in")
    vis = (f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{dur - 1 if out else 0}"/>'
           f'</p:stCondLst></p:cTn>{tgt(s)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
           f'</p:cBhvr><p:to><p:strVal val="{"hidden" if out else "visible"}"/></p:to></p:set>')
    eff = (f'<p:animEffect transition="{tr}" filter="fade"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}"/>{tgt(s)}'
           f'</p:cBhvr></p:animEffect>')
    return (f'<p:par><p:cTn id="{nid()}" presetID="10" presetClass="{cls}" presetSubtype="0" fill="hold" '
            f'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + ((eff + vis) if out else (vis + eff)) + '</p:childTnLst></p:cTn></p:par>')


def move(s, x0, x1, node="withEffect", dur=1200):
    return (f'<p:par><p:cTn id="{nid()}" presetID="0" presetClass="path" presetSubtype="0" accel="30000" '
            f'decel="30000" fill="hold" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
            f'<p:childTnLst><p:animMotion origin="layout" path="M {x0:.5f} 0 L {x1:.5f} 0 E" pathEditMode="relative" '
            f'rAng="0" ptsTypes="AA"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}" fill="hold"/>{tgt(s)}<p:attrNameLst>'
            f'<p:attrName>ppt_x</p:attrName><p:attrName>ppt_y</p:attrName></p:attrNameLst></p:cBhvr>'
            f'<p:rCtr x="0" y="0"/></p:animMotion></p:childTnLst></p:cTn></p:par>')


def par(delay, effects):
    return (f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
            f'<p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par>')


def first(effects, node="clickEffect"):
    effects[0] = effects[0].replace('nodeType="withEffect"', f'nodeType="{node}"', 1)
    return effects


def open_stop(st):
    return [fade(st["dim"], True), fade(st["lit"]), fade(st["card"]), fade(st["body"])]


frac = lambda px: px / 1600
pos = [110] + XS + [1450]
main = []
for k, st in enumerate(stops):
    main.append(par("indefinite", [par(0, first([move(vial, frac(pos[k] - 110), frac(pos[k + 1] - 110))])),
                                   par(1200, first(open_stop(st), "afterEffect"))]))
main.append(par("indefinite", [par(0, first([move(vial, frac(pos[3] - 110), frac(pos[4] - 110)), fade(banner)]))]))

triggers = []
for st in stops:
    for hit in (st["base"], st["pic"]):
        on = f'<p:cond evt="onClick" delay="0">{tgt(hit)}</p:cond>'
        triggers.append(
            f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{nid()}" restart="whenNotActive" fill="hold" '
            f'evtFilter="cancelBubble" nodeType="interactiveSeq"><p:stCondLst>{on}</p:stCondLst>'
            f'<p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync><p:childTnLst>'
            + par(0, [par(0, first(open_stop(st)))]) +
            f'</p:childTnLst></p:cTn><p:nextCondLst>{on}</p:nextCondLst></p:seq>')

timing = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
          '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
          '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
          + "".join(main) + '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl>'
          '<p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl>'
          '<p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>' + "".join(triggers)
          + '</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
slide._element.append(etree.fromstring(timing))
slide.notes_slide.notes_text_frame.text = (
    "About 1 min. The plan for today, as a vial's trip along the line; each click moves it to the next stop. "
    "Clicking a stop's circle or picture opens it out of order, for questions.\n"
    "1. Describe it: what a module is (its parts, its IEC 61499 program, its AAS), how products, processes and "
    "resources are modelled, and the rules that turn a skill in the AAS into function blocks.\n"
    "2. Change it: a product asks for something the line cannot do yet; a new skill is made from the module's "
    "primitives, and modsync builds it while the module runs, checks it and records it.\n"
    "3. Let it decide: the next step, agents. Products call for proposals, modules bid, and a module's agent "
    "reconfigures its own module with the same tools.\n"
    "Last click: the one thread through all three.")
prs.save(OUT)
print("wrote", OUT)
