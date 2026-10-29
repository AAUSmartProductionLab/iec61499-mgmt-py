"""One explainer slide for IEC 61499: what a function block is, event and data connections, WITH, and a
blinking network (E_CYCLE → E_T_FF → QX, an LED on GPIO 17) whose connections light up as it runs.
Native PowerPoint shapes, in the deck's style, so the slide can be pasted into any deck and edited.

    python iec61499_explainer.py <out.pptx>          (python-pptx and lxml)
"""
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

PX = 7620                                        # the slide as a 1600 × 900 grid
TEAL, TEAL_D, INK, GREY, LINE = "2A918B", "0F726C", "2D3134", "5C6166", "BEC2CA"
EVENT, DATA = "5B963B", "4D7392"                 # event connections green, data connections blue
EVENT_HI, DATA_HI, FLASH = "8CCB5E", "7FB0DD", "F07D2E"
BLOCK_FILL, BLOCK_LINE, TYPE = "F4F9FC", "41484B", "4D7392"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
slide = prs.slides.add_slide(prs.slide_layouts[6])
S = slide.shapes
E = lambda v: Emu(int(round(v * PX)))


def rgb(h):
    return RGBColor.from_string(h)


def textbox(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, name=None):
    """runs: [(text, pt, bold, colour, italic)], one paragraph; or a list of such lists, one per paragraph."""
    tb = S.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paras = runs if runs and isinstance(runs[0], list) else [runs]
    for k, pr in enumerate(paras):
        p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
        p.alignment = align
        for text, pt, bold, colour, *it in pr:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name, f.size, f.bold, f.color.rgb = "Arial", Pt(pt), bold, rgb(colour)
            f.italic = bool(it and it[0])
    if name:
        tb.name = name
    return tb


def rect(x, y, w, h, fill=None, line=None, width=1.25, shape=MSO_SHAPE.RECTANGLE, name=None):
    sp = S.add_shape(shape, E(x), E(y), E(w), E(h))
    if fill:
        sp.fill.solid(), setattr(sp.fill.fore_color, "rgb", rgb(fill))
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb, sp.line.width = rgb(line), Pt(width)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    if name:
        sp.name = name
    return sp


def polyline(pts, colour, width, name=None, dash=False):
    fb = S.build_freeform(E(pts[0][0]), E(pts[0][1]), scale=1.0)
    fb.add_line_segments([(E(x), E(y)) for x, y in pts[1:]], close=False)
    sp = fb.convert_to_shape()
    sp.fill.background()
    sp.line.color.rgb, sp.line.width = rgb(colour), Pt(width)
    if dash:
        sp.line.dash_style = 4                   # dash
    sp.shadow.inherit = False
    if name:
        sp.name = name
    return sp


def pin(x, y, colour, right=True):
    t = rect(x - 1 if right else x - 9, y - 5, 10, 10, colour, None, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    t.rotation = 90 if right else 270


# ---- title -------------------------------------------------------------------------------------------
textbox(85, 40, 900, 24, [("IEC 61499", 12, True, TEAL_D)])
textbox(85, 70, 1430, 60, [("IEC 61499: events say when, data says what", 28, True, TEAL)])

# ---- the function blocks -----------------------------------------------------------------------------
PIN, NECK = 32, 34


def fb(name, typ, x, y, w, ev_in=(), ev_out=(), d_in=(), d_out=()):
    ne, nd = max(len(ev_in), len(ev_out), 1), max(len(d_in), len(d_out), 1)
    hh, bh = 12 + PIN * ne, 12 + PIN * nd
    yb = y + hh + NECK
    shapes = [rect(x, y, w, hh, BLOCK_FILL, BLOCK_LINE), rect(x + 16, y + hh, w - 32, NECK, BLOCK_FILL, BLOCK_LINE),
              rect(x, yb, w, bh, BLOCK_FILL, BLOCK_LINE)]
    textbox(x - 30, y - 30, w + 60, 26, [(name, 14, True, INK)], PP_ALIGN.CENTER)
    textbox(x + 16, y + hh, w - 32, NECK, [(typ, 12, False, TYPE, True)], PP_ALIGN.CENTER)
    pins = {}
    for i, e in enumerate(ev_in):
        py = y + 6 + PIN * i + PIN / 2
        pin(x + 1, py, EVENT), textbox(x + 14, py - 12, w / 2, 24, [(e, 12, False, EVENT)])
        pins[("in", e)] = (x, py)
    for i, e in enumerate(ev_out):
        py = y + 6 + PIN * i + PIN / 2
        textbox(x + w / 2 - 14, py - 12, w / 2, 24, [(e, 12, False, EVENT)], PP_ALIGN.RIGHT), pin(x + w - 9, py, EVENT)
        pins[("out", e)] = (x + w, py)
    for i, d in enumerate(d_in):
        py = yb + 6 + PIN * i + PIN / 2
        pin(x + 1, py, DATA), textbox(x + 14, py - 12, w / 2, 24, [(d, 12, False, DATA)])
        pins[("in", d)] = (x, py)
    for i, d in enumerate(d_out):
        py = yb + 6 + PIN * i + PIN / 2
        textbox(x + w / 2 - 14, py - 12, w / 2, 24, [(d, 12, False, DATA)], PP_ALIGN.RIGHT), pin(x + w - 9, py, DATA)
        pins[("out", d)] = (x + w, py)
    return pins, (x, y - 32, w, hh + NECK + bh + 32)


Y = 230
START, b_start = fb("START", "E_RESTART", 80, Y, 170, ev_out=["COLD", "WARM", "STOP"])
CYC, b_cyc = fb("Cycle", "E_CYCLE", 420, Y, 180, ev_in=["START", "STOP"], ev_out=["EO"], d_in=["DT"])
TFF, b_tff = fb("Toggle", "E_T_FF", 760, Y, 170, ev_in=["CLK"], ev_out=["EO"], d_out=["Q"])
LED, b_led = fb("LED", "QX", 1090, Y, 230, ev_in=["INIT", "REQ"], ev_out=["INITO", "CNF"],
                d_in=["QI", "PARAMS", "OUT"], d_out=["QO", "STATUS"])

# the lamp on GPIO 17
polyline([(1320, 300), (1372, 300)], GREY, 1.5, dash=True)
bulb_off = rect(1372, 258, 84, 84, "D9DCDF", "9AA0A4", 1.5, MSO_SHAPE.OVAL, "Lamp off")
bulb_on = rect(1372, 258, 84, 84, "F6C343", "D99A1E", 2.5, MSO_SHAPE.OVAL, "Lamp on")
textbox(1340, 350, 150, 24, [("GPIO 17", 12, True, GREY)], PP_ALIGN.CENTER)

# ---- connections, constants, WITH --------------------------------------------------------------------
P = lambda pins, io, n: pins[(io, n)]
ev_routes = {
    "COLD to Cycle.START": [P(START, "out", "COLD"), P(CYC, "in", "START")],
    "COLD to LED.INIT": [P(START, "out", "COLD"), (290, 252), (290, 190), (1060, 190), (1060, 252), P(LED, "in", "INIT")],
    "Cycle.EO to Toggle.CLK": [P(CYC, "out", "EO"), P(TFF, "in", "CLK")],
    "Toggle.EO to LED.REQ": [P(TFF, "out", "EO"), (1010, 252), (1010, 284), P(LED, "in", "REQ")],
}
data_routes = {"Toggle.Q to LED.OUT": [P(TFF, "out", "Q"), (960, 330), (960, 426), P(LED, "in", "OUT")]}
ev_lines = [polyline(pts, EVENT, 2.25, f"Event {k}") for k, pts in ev_routes.items()]
data_lines = [polyline(pts, DATA, 2.25, f"Data {k}") for k, pts in data_routes.items()]
consts = [textbox(290, P(CYC, "in", "DT")[1] - 12, 104, 24, [("T#1s", 12, True, INK)], PP_ALIGN.RIGHT),
          textbox(966, P(LED, "in", "QI")[1] - 12, 92, 24, [("TRUE", 12, True, INK)], PP_ALIGN.RIGHT),
          textbox(966, P(LED, "in", "PARAMS")[1] - 12, 92, 24, [("'17'", 12, True, INK)], PP_ALIGN.RIGHT)]


def with_bar(x, ys, name):
    """The WITH notation: a square on each associated pin, joined by a line."""
    out = [polyline([(x, min(ys)), (x, max(ys))], BLOCK_LINE, 1.5)]
    out += [rect(x - 5, y - 5, 10, 10, "FFFFFF", BLOCK_LINE, 1.5) for y in ys]
    for s in out:
        s.name = name
    return out


withs = (with_bar(410, [P(CYC, "in", "START")[1], P(CYC, "in", "DT")[1]], "WITH START DT")
         + with_bar(942, [P(TFF, "out", "EO")[1], P(TFF, "out", "Q")[1]], "WITH EO Q")
         + with_bar(1080, [P(LED, "in", "REQ")[1], P(LED, "in", "OUT")[1]], "WITH REQ OUT")
         + with_bar(1072, [P(LED, "in", "INIT")[1], P(LED, "in", "QI")[1], P(LED, "in", "PARAMS")[1]], "WITH INIT QI PARAMS"))

# ---- what lights up while it runs --------------------------------------------------------------------
hi = {k: polyline(pts, EVENT_HI, 6, f"Lit {k}") for k, pts in ev_routes.items()}
hi.update({k: polyline(pts, DATA_HI, 6, f"Lit {k}") for k, pts in data_routes.items()})
flash = {n: rect(b[0] - 10, b[1] - 4, b[2] + 20, b[3] + 14, None, FLASH, 3.5, MSO_SHAPE.ROUNDED_RECTANGLE, f"Runs {n}")
         for n, b in [("Start", b_start), ("Cycle", b_cyc), ("Toggle", b_tff), ("LED", b_led)]}
for f in flash.values():
    f.adjustments[0] = 0.08
q_false = rect(826, 362, 118, 30, "FFFFFF", DATA, 1.5, MSO_SHAPE.ROUNDED_RECTANGLE, "Q is FALSE")
q_false.text_frame.text = ""
textbox(826, 362, 118, 30, [("Q = FALSE", 12, True, DATA)], PP_ALIGN.CENTER)
q_true = rect(826, 362, 118, 30, "E1EAF2", DATA, 1.5, MSO_SHAPE.ROUNDED_RECTANGLE, "Q is TRUE")
tf = q_true.text_frame
tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Q = TRUE"
r.font.name, r.font.size, r.font.bold, r.font.color.rgb = "Arial", Pt(12), True, rgb(DATA)

# ---- click 1: what a block is ---------------------------------------------------------------------------
anat = [rect(b_tff[0] - 10, b_tff[1] - 4, b_tff[2] + 20, b_tff[3] + 14, None, TEAL, 3,
             MSO_SHAPE.ROUNDED_RECTANGLE, "Block highlighted"),
        textbox(700, 140, 300, 40, [[("Events", 13, True, EVENT), (": when it runs", 13, False, INK)]],
                PP_ALIGN.CENTER),
        textbox(700, 404, 300, 40, [[("Data", 13, True, DATA), (": what it works on", 13, False, INK)]],
                PP_ALIGN.CENTER)]
anat[0].adjustments[0] = 0.08

# ---- the cards below -------------------------------------------------------------------------------------
CARDS = [("A function block", None,
          "Events on top, data below, its type in between. It runs only when an event arrives: no event, nothing happens."),
         ("Event connection", EVENT,
          "When. A pulse from an event output starts the next block: the order of the network."),
         ("Data connection", DATA,
          "What. A value from a data output, or a constant set at the pin, such as T#1s or '17'."),
         ("WITH", "with",
          "Which data goes with which event: Q leaves with EO, OUT is read when REQ arrives. Nothing in between.")]
cards = []
for i, (title, kind, body) in enumerate(CARDS):
    x, y, w, h = 80 + i * 365, 560, 335, 250
    group = [rect(x, y, w, h, "F7F8F9", "D9DCDF", 1, MSO_SHAPE.ROUNDED_RECTANGLE, f"Card {title}")]
    group[0].adjustments[0] = 0.06
    if kind in (EVENT, DATA):
        group.append(polyline([(x + 22, y + 37), (x + 82, y + 37)], kind, 3))
    elif kind == "with":
        group += with_bar(x + 40, [y + 24, y + 48], "Card WITH")
    else:
        group.append(rect(x + 22, y + 22, 34, 14, BLOCK_FILL, BLOCK_LINE, 1.25))
        group.append(rect(x + 28, y + 36, 22, 8, BLOCK_FILL, BLOCK_LINE, 1.25))
        group.append(rect(x + 22, y + 44, 34, 12, BLOCK_FILL, BLOCK_LINE, 1.25))
    group.append(textbox(x + 100, y + 20, w - 110, 34, [(title, 15, True, INK)]))
    group.append(textbox(x + 22, y + 78, w - 44, h - 96, [(body, 13, False, GREY)], anchor=MSO_ANCHOR.TOP))
    cards.append(group)
run_note = textbox(1350, 400, 220, 140, [[("Running", 14, True, FLASH)],
                                          [("E_CYCLE fires EO every second, E_T_FF toggles Q, QX writes Q to GPIO 17:"
                                            " the lamp blinks.", 12, False, INK)]], anchor=MSO_ANCHOR.TOP)

# ---- the animation ---------------------------------------------------------------------------------------
spid = lambda s: s.shape_id
_id = [2]


def nid():
    _id[0] += 1
    return _id[0]


def tgt(s):
    return f'<p:tgtEl><p:spTgt spid="{spid(s)}"/></p:tgtEl>'


def appear(s, node="withEffect", delay=0):
    return (f'<p:par><p:cTn id="{nid()}" presetID="1" presetClass="entr" presetSubtype="0" fill="hold" '
            f'nodeType="{node}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
            f'</p:stCondLst></p:cTn>{tgt(s)}<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
            f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')


def fade(s, out=False, node="withEffect", dur=400):
    cls, tr = ("exit", "out") if out else ("entr", "in")
    vis = (f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="{dur - 1 if out else 0}"/>'
           f'</p:stCondLst></p:cTn>{tgt(s)}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
           f'</p:cBhvr><p:to><p:strVal val="{"hidden" if out else "visible"}"/></p:to></p:set>')
    eff = (f'<p:animEffect transition="{tr}" filter="fade"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}"/>{tgt(s)}'
           f'</p:cBhvr></p:animEffect>')
    return (f'<p:par><p:cTn id="{nid()}" presetID="10" presetClass="{cls}" presetSubtype="0" fill="hold" '
            f'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            + ((eff + vis) if out else (vis + eff)) + '</p:childTnLst></p:cTn></p:par>')


def wipe(s, node="withEffect", dur=700):
    return (f'<p:par><p:cTn id="{nid()}" presetID="22" presetClass="entr" presetSubtype="8" fill="hold" '
            f'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:set><p:cBhvr><p:cTn id="{nid()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
            f'</p:stCondLst></p:cTn>{tgt(s)}<p:attrNameLst><p:attrName>style.visibility</p:attrName>'
            f'</p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
            f'<p:animEffect transition="in" filter="wipe(right)"><p:cBhvr><p:cTn id="{nid()}" dur="{dur}"/>{tgt(s)}'
            f'</p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')


PERIOD = 2000                                    # two ticks of E_CYCLE, one on and one off of the lamp


def blink(s, keys, node="withEffect"):
    """Shown and hidden at fixed times within PERIOD, again and again until the slide ends (Blink, by keys)."""
    tav = "".join(f'<p:tav tm="{int(t * 100000 / PERIOD)}"><p:val><p:strVal val="{v}"/></p:val></p:tav>'
                  for t, v in keys)
    return (f'<p:par><p:cTn id="{nid()}" presetID="35" presetClass="emph" presetSubtype="0" fill="hold" '
            f'nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
            f'<p:anim calcmode="discrete" valueType="str"><p:cBhvr override="childStyle"><p:cTn id="{nid()}" '
            f'dur="{PERIOD}" repeatCount="indefinite" fill="hold"/>{tgt(s)}<p:attrNameLst>'
            f'<p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst>{tav}</p:tavLst>'
            f'</p:anim></p:childTnLst></p:cTn></p:par>')


def group(delay, effects):
    return (f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
            f'<p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par>')


def click(*groups):
    return (f'<p:par><p:cTn id="{nid()}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
            f'<p:childTnLst>{"".join(groups)}</p:childTnLst></p:cTn></p:par>')


def first(effects):
    """The first effect of a click is its click effect."""
    effects[0] = effects[0].replace('nodeType="withEffect"', 'nodeType="clickEffect"', 1)
    return effects


def after(effects):
    effects[0] = effects[0].replace('nodeType="withEffect"', 'nodeType="afterEffect"', 1)
    return effects


ON, OFF = "visible", "hidden"
loop = [
    blink(flash["Cycle"], [(0, ON), (150, OFF), (1000, ON), (1150, OFF)]),
    blink(hi["Cycle.EO to Toggle.CLK"], [(0, ON), (300, OFF), (1000, ON), (1300, OFF)]),
    blink(flash["Toggle"], [(0, OFF), (300, ON), (500, OFF), (1300, ON), (1500, OFF)]),
    blink(q_true, [(0, OFF), (400, ON), (1400, OFF)]),
    blink(hi["Toggle.EO to LED.REQ"], [(0, OFF), (400, ON), (700, OFF), (1400, ON), (1700, OFF)]),
    blink(hi["Toggle.Q to LED.OUT"], [(0, OFF), (400, ON), (700, OFF), (1400, ON), (1700, OFF)]),
    blink(flash["LED"], [(0, OFF), (700, ON), (900, OFF), (1700, ON), (1900, OFF)]),
    blink(bulb_on, [(0, OFF), (700, ON), (1700, OFF)]),
]
sequence = [
    click(group(0, first([fade(s) for s in anat] + [fade(s) for s in cards[0]]))),
    click(group(0, first([fade(s, True) for s in anat] + [wipe(s) for s in ev_lines] + [fade(s) for s in cards[1]]))),
    click(group(0, first([wipe(s) for s in data_lines] + [fade(s) for s in consts] + [fade(s) for s in cards[2]]))),
    click(group(0, first([fade(s) for s in withs] + [fade(s) for s in cards[3]]))),
    click(group(0, first([fade(run_note), appear(flash["Start"]), appear(hi["COLD to Cycle.START"]),
                          appear(hi["COLD to LED.INIT"])])),
          group(500, after([fade(flash["Start"], True, dur=200), fade(hi["COLD to Cycle.START"], True, dur=200),
                            fade(hi["COLD to LED.INIT"], True, dur=200)])),
          group(700, after([appear(flash["Cycle"]), appear(flash["Toggle"]), appear(flash["LED"]),
                      appear(hi["Cycle.EO to Toggle.CLK"]), appear(hi["Toggle.EO to LED.REQ"]),
                      appear(hi["Toggle.Q to LED.OUT"]), appear(q_true), appear(bulb_on)] + loop))),
]
timing = ('<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
          '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
          '<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
          + "".join(sequence) + '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl>'
          '<p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl>'
          '<p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
          '</p:timing>')
sld = slide._element
sld.append(etree.fromstring(timing))

slide.notes_slide.notes_text_frame.text = (
    "About 2 min, five clicks. IEC 61499 in one picture: a small network that blinks a lamp.\n"
    "1. A function block: event inputs and outputs on top, data inputs and outputs below, its type in the "
    "middle. It runs only when an event arrives, unlike a PLC scan.\n"
    "2. Event connections (green) say when: START's COLD starts the cycle and initialises the output; the "
    "cycle's EO clocks the toggle; the toggle's EO asks the output block to write.\n"
    "3. Data connections (blue) say what: the toggle's Q goes to the output's OUT; constants sit at the pins: "
    "DT = T#1s, QI = TRUE, PARAMS = '17', the GPIO pin.\n"
    "4. WITH ties data to events: Q is sent with EO, OUT is read when REQ arrives; a block reads its inputs only "
    "then, so values cannot change under it.\n"
    "5. Running: once a second E_CYCLE fires, E_T_FF flips Q, QX writes it to GPIO 17, and the lamp blinks; the "
    "orange frames show which block runs. In 4diac this is the classic Blinky.")
prs.save(sys.argv[1])
print("wrote", sys.argv[1])
